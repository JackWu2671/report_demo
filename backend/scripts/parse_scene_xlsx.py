#!/usr/bin/env python3
"""
parse_scene_xlsx.py — 将 场景.xlsx 转换为 scene.json

用法:
  python3 parse_scene_xlsx.py

输入: expert_knowledge/场景.xlsx
输出: expert_knowledge/scene.json

字段说明:
  id          — 场景唯一标识
  name        — 场景名称（同 SCENEKEY）
  level       — 固定为 "场景"
  description — 一句话描述
  detail      — 详细拓展逻辑（原字段名 expandLogic）
  keywords    — 关键词列表（原字段名 keyWords）
  sampleIssue — 示例提问
  condition   — 触发条件（原数据无此字段，默认空字符串）
  dimensions  — 子维度列表，每项只保留 id / name / rank
"""

import json
import os
import sys

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_BACKEND_DIR = os.path.dirname(_SCRIPT_DIR)
_KB_DIR = os.path.join(_BACKEND_DIR, "expert_knowledge")

INPUT_FILE  = os.path.join(_KB_DIR, "场景.xlsx")
OUTPUT_FILE = os.path.join(_KB_DIR, "scene.json")


def parse_dimensions(raw) -> list[dict]:
    """只保留每个 dimension 的 id / name / rank。"""
    if not isinstance(raw, list):
        return []
    result = []
    for d in raw:
        result.append({
            "id":   d.get("id",   ""),
            "name": d.get("name", ""),
            "rank": d.get("rank", None),
        })
    # 按 rank 升序排列
    result.sort(key=lambda x: (x["rank"] is None, x["rank"]))
    return result


def parse_content(content_str: str) -> dict | None:
    """解析 CONTENT 列的 JSON 字符串，失败返回 None。"""
    try:
        return json.loads(content_str)
    except (json.JSONDecodeError, TypeError):
        return None


def convert_row(scene_key: str, content_str: str) -> dict | None:
    """将一行 Excel 数据转换为目标结构。"""
    obj = parse_content(content_str)
    if obj is None:
        print(f"[警告] SCENEKEY={scene_key!r} 的 CONTENT 解析失败，跳过", file=sys.stderr)
        return None

    return {
        "id":          obj.get("id", ""),
        "name":        obj.get("name", scene_key),
        "level":       "场景",
        "description": obj.get("description", ""),
        "detail":      obj.get("expandLogic", ""),
        "keywords":    obj.get("keyWords", []),
        "sampleIssue": obj.get("sampleIssue", ""),
        "condition":   "",
        "dimensions":  parse_dimensions(obj.get("dimensions", [])),
    }


def main():
    try:
        import openpyxl
    except ImportError:
        print("[错误] 请先安装 openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    if not os.path.exists(INPUT_FILE):
        print(f"[错误] 文件不存在: {INPUT_FILE}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    ws = wb.active

    # 读取表头，找 SCENEKEY / CONTENT 列索引（大小写不敏感）
    headers = [str(cell.value).strip().upper() if cell.value else "" for cell in ws[1]]
    try:
        idx_key     = headers.index("SCENEKEY")
        idx_content = headers.index("CONTENT")
    except ValueError:
        print(f"[错误] 找不到必要列，实际表头: {headers}", file=sys.stderr)
        sys.exit(1)

    scenes = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        scene_key    = str(row[idx_key]).strip()   if row[idx_key]     else ""
        content_str  = str(row[idx_content]).strip() if row[idx_content] else ""

        if not scene_key and not content_str:
            continue  # 跳过空行

        item = convert_row(scene_key, content_str)
        if item:
            scenes.append(item)

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(scenes, f, ensure_ascii=False, indent=2)

    print(f"转换完成，共 {len(scenes)} 条场景 → {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
