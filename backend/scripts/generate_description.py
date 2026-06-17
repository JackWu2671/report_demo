#!/usr/bin/env python3
"""
generate_description.py — 批量为评估项.xlsx 生成 DESCRIPTION 列

读取 reference/评估项.xlsx，对每行 CONTENT.expandLogic 调用 LLM
生成 description，写入 DESCRIPTION 列（已有值的行跳过），覆盖前自动备份。

用法（在 backend/ 目录下运行）：
  cd backend
  python3 scripts/generate_description.py

依赖环境变量（同 api_server.py）：
  LLM_BASE_URL, LLM_MODEL_NAME, LLM_API_KEY
"""

import asyncio
import json
import logging
import shutil
import sys
from pathlib import Path

# sys.path 调整和环境变量加载属于导入准备，放在本地模块导入之前
sys.path.insert(0, str(Path(__file__).parent.parent))

from dotenv import load_dotenv
load_dotenv(Path(__file__).parent.parent / ".env")

import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from llm.config import LLMConfig
from services.llm_service import LLMService

# ── 常量（所有 import 完成后统一声明）────────────────────────────────
logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

_SCRIPT_DIR = Path(__file__).parent
_BACKEND_DIR = _SCRIPT_DIR.parent
_KB_DIR = _BACKEND_DIR / "reference"
INPUT_FILE = _KB_DIR / "评估项.xlsx"
PROMPT_FILE = _SCRIPT_DIR / "generate_description_prompt.txt"
CONCURRENCY = 3       # 并发调用数，避免触发限流
SKIP_NONEMPTY = True  # True=跳过已有 DESCRIPTION 的行，False=全量重新生成
# ─────────────────────────────────────────────────────────────────────


def _load_prompt() -> str:
    with open(PROMPT_FILE, encoding="utf-8") as f:
        return f.read()


def _get_expand_logic(content_str: str) -> str:
    try:
        obj = json.loads(content_str)
        return obj.get("expandLogic", "")
    except (json.JSONDecodeError, TypeError):
        return ""


async def _generate_one(llm: LLMService, prompt_template: str, expand_logic: str) -> str:
    prompt = prompt_template.replace("{expand_logic}", expand_logic)
    messages = [{"role": "user", "content": prompt}]
    config = LLMConfig(max_tokens=512, temperature=0.3)
    result = await llm.complete(messages, config)
    return result.strip()


def _validate_inputs() -> None:
    for path, label in [(INPUT_FILE, "评估项.xlsx"), (PROMPT_FILE, "prompt 文件")]:
        if not path.exists():
            logging.error("%s 不存在: %s", label, path)
            sys.exit(1)


def _resolve_columns(ws: Worksheet) -> tuple[int, int | None, int]:
    """返回 (idx_content, idx_key, idx_desc)，列号均为 1-based。"""
    headers = [str(c.value).strip().upper() if c.value else "" for c in ws[1]]

    if "CONTENT" not in headers:
        logging.error("找不到 CONTENT 列")
        sys.exit(1)

    idx_content = headers.index("CONTENT") + 1
    idx_key = (headers.index("SCENEKEY") + 1) if "SCENEKEY" in headers else None

    if "DESCRIPTION" in headers:
        idx_desc = headers.index("DESCRIPTION") + 1
    else:
        idx_desc = len(headers) + 1
        ws.cell(row=1, column=idx_desc, value="DESCRIPTION")
        logging.info("已新增 DESCRIPTION 列（第 %d 列）", idx_desc)

    return idx_content, idx_key, idx_desc


def _collect_rows(
    ws: Worksheet,
    idx_content: int,
    idx_desc: int,
    idx_key: int | None,
) -> list[tuple[int, str, str]]:
    """返回需要生成 description 的行列表：[(row_idx, name, expand_logic)]。"""
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        content_val = ws.cell(row=row_idx, column=idx_content).value
        desc_val = ws.cell(row=row_idx, column=idx_desc).value
        name = str(ws.cell(row=row_idx, column=idx_key).value or "").strip() \
            if idx_key else f"行{row_idx}"

        content_str = str(content_val).strip() if content_val else ""
        if not content_str:
            continue
        if SKIP_NONEMPTY and desc_val and str(desc_val).strip():
            logging.info("跳过 %s（已有 description）", name)
            continue
        expand_logic = _get_expand_logic(content_str)
        if not expand_logic:
            logging.info("跳过 %s（expandLogic 为空）", name)
            continue
        rows.append((row_idx, name, expand_logic))
    return rows


async def _run_batch(
    llm: LLMService,
    prompt_template: str,
    ws: Worksheet,
    rows: list[tuple[int, str, str]],
    idx_desc: int,
) -> tuple[int, int]:
    """并发调用 LLM，返回 (done, failed)。"""
    total = len(rows)
    sem = asyncio.Semaphore(CONCURRENCY)
    done = 0
    failed = 0

    async def _process(row_idx: int, name: str, expand_logic: str) -> None:
        nonlocal done, failed
        async with sem:
            try:
                desc = await _generate_one(llm, prompt_template, expand_logic)
                ws.cell(row=row_idx, column=idx_desc, value=desc)
                done += 1
                preview = desc[:80] + ("…" if len(desc) > 80 else "")
                logging.info("[%d/%d] ✓ %s\n  %s", done + failed, total, name, preview)
            except Exception as exc:
                failed += 1
                logging.error("[%d/%d] ✗ %s: %s", done + failed, total, name, exc)

    await asyncio.gather(*[_process(r, n, e) for r, n, e in rows])
    return done, failed


def _backup_and_save(wb: Workbook) -> None:
    backup = INPUT_FILE.with_suffix(".bak.xlsx")
    shutil.copy(INPUT_FILE, backup)
    logging.info("原文件已备份至 %s", backup.name)
    wb.save(INPUT_FILE)


async def main() -> None:
    _validate_inputs()

    prompt_template = _load_prompt()
    llm = LLMService.from_env()

    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active

    idx_content, idx_key, idx_desc = _resolve_columns(ws)
    rows = _collect_rows(ws, idx_content, idx_desc, idx_key)

    logging.info("共需生成 %d 条", len(rows))
    if not rows:
        logging.info("无需处理，退出。")
        return

    done, failed = await _run_batch(llm, prompt_template, ws, rows, idx_desc)

    _backup_and_save(wb)
    logging.info("完成：成功 %d 条，失败 %d 条 → %s", done, failed, INPUT_FILE)


if __name__ == "__main__":
    asyncio.run(main())
